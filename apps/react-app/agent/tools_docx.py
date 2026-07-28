"""LangChain tool: create a downloadable Word (.docx) document.

Visual styling is loaded from a skill's ``references/theme.yml`` (default:
``create-docx-report``), not hardcoded in this module. Content structure
guidance lives in that skill's ``SKILL.md``.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Any, Optional

import yaml
from langchain_core.tools import StructuredTool

from server.artifacts import DOCX_MEDIA_TYPE, save_artifact

logger = logging.getLogger(__name__)

_APP_ROOT = Path(__file__).resolve().parent.parent
_SKILLS_DIR = _APP_ROOT / "skills"
_DEFAULT_SKILL = "create-docx-report"

_HEADING_RE = re.compile(r"^(#{1,3})\s+(.+)$")
_BULLET_RE = re.compile(r"^[-*•]\s+(.+)$")
_NUMBERED_RE = re.compile(r"^\d+[.)]\s+(.+)$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
_HR_RE = re.compile(r"^-{3,}$|^\*{3,}$|^_{3,}$")

_FALLBACK_THEME: dict[str, Any] = {
    "colors": {
        "accent": "#D52B1E",
        "white": "#FFFFFF",
        "near_black": "#1A1A1A",
        "muted_gray": "#666666",
        "alt_row": "#F7F7F7",
        "soft_border": "#E5E5E5",
    },
    "typography": {
        "body_font": "Calibri",
        "heading_font": "Calibri",
        "title_size_pt": 22,
        "h1_size_pt": 18,
        "h2_size_pt": 14,
        "h3_size_pt": 12,
        "body_size_pt": 11,
        "table_header_size_pt": 10,
        "table_body_size_pt": 10,
        "notes_size_pt": 9,
    },
    "layout": {
        "margin_inches": 1.0,
        "line_spacing": 1.15,
        "title_rule": True,
    },
}


def _hex_to_rgb(value: str) -> tuple[int, int, int]:
    h = (value or "").strip().lstrip("#")
    if len(h) != 6:
        raise ValueError(f"Invalid hex color: {value!r}")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _deep_merge(base: dict, override: dict) -> dict:
    out = dict(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _deep_merge(out[k], v)
        else:
            out[k] = v
    return out


def load_skill_theme(skill_name: str = _DEFAULT_SKILL) -> dict[str, Any]:
    """Load theme tokens from ``skills/<skill>/references/theme.yml``."""
    name = (skill_name or _DEFAULT_SKILL).strip() or _DEFAULT_SKILL
    # Only allow simple skill folder names (no path traversal).
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", name):
        logger.warning("Invalid skill_name %r — using %s", skill_name, _DEFAULT_SKILL)
        name = _DEFAULT_SKILL

    theme_path = _SKILLS_DIR / name / "references" / "theme.yml"
    if not theme_path.is_file():
        logger.warning("Theme not found at %s — using fallback", theme_path)
        return dict(_FALLBACK_THEME)

    try:
        raw = yaml.safe_load(theme_path.read_text(encoding="utf-8")) or {}
        theme = _deep_merge(_FALLBACK_THEME, raw if isinstance(raw, dict) else {})
        logger.info("Loaded DOCX theme from skill '%s' (%s)", name, theme_path)
        return theme
    except Exception:
        logger.exception("Failed to load theme from %s — using fallback", theme_path)
        return dict(_FALLBACK_THEME)


def _theme_colors(theme: dict) -> dict[str, tuple[int, int, int]]:
    colors = theme.get("colors") or {}
    out = {}
    for key, default_hex in _FALLBACK_THEME["colors"].items():
        try:
            out[key] = _hex_to_rgb(colors.get(key) or default_hex)
        except ValueError:
            out[key] = _hex_to_rgb(default_hex)
    return out


def _theme_type(theme: dict) -> dict[str, Any]:
    return {**_FALLBACK_THEME["typography"], **(theme.get("typography") or {})}


def _theme_layout(theme: dict) -> dict[str, Any]:
    return {**_FALLBACK_THEME["layout"], **(theme.get("layout") or {})}


def _rgb(rgb_tuple):
    from docx.shared import RGBColor

    return RGBColor(*rgb_tuple)


def _set_run_font(run, *, size_pt=11, bold=False, color=(0x1A, 0x1A, 0x1A), font="Calibri"):
    from docx.shared import Pt

    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _style_paragraph(paragraph, *, space_after_pt=8, space_before_pt=0, line_spacing=1.15):
    from docx.shared import Pt

    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after_pt)
    pf.space_before = Pt(space_before_pt)
    pf.line_spacing = line_spacing


def _shade_cell(cell, rgb_tuple):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:shd"))
    if existing is not None:
        tc_pr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "%02X%02X%02X" % rgb_tuple)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_cell_borders(cell, color_hex="E5E5E5", sz="4"):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color_hex.lstrip("#").upper())
        borders.append(el)
    tc_pr.append(borders)


def _apply_document_defaults(doc, theme: dict) -> None:
    from docx.shared import Inches, Pt
    from docx.oxml.ns import qn

    colors = _theme_colors(theme)
    typo = _theme_type(theme)
    layout = _theme_layout(theme)
    margin = float(layout.get("margin_inches", 1.0))
    body_font = typo["body_font"]
    heading_font = typo["heading_font"]
    accent = colors["accent"]
    near_black = colors["near_black"]

    for section in doc.sections:
        section.top_margin = Inches(margin)
        section.bottom_margin = Inches(margin)
        section.left_margin = Inches(margin)
        section.right_margin = Inches(margin)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(int(typo["body_size_pt"]))
    normal.font.color.rgb = _rgb(near_black)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), body_font)
    rfonts.set(qn("w:hAnsi"), body_font)

    for level, key in ((1, "title_size_pt"), (2, "h2_size_pt"), (3, "h3_size_pt")):
        try:
            hs = styles[f"Heading {level}"]
        except KeyError:
            continue
        hs.font.name = heading_font
        hs.font.size = Pt(int(typo[key]))
        hs.font.bold = True
        hs.font.color.rgb = _rgb(accent)
        hs.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        hs.paragraph_format.space_after = Pt(8)


def _add_brand_rule(doc, theme: dict) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    accent_hex = (theme.get("colors") or {}).get("accent") or "#D52B1E"
    layout = _theme_layout(theme)
    line_spacing = float(layout.get("line_spacing", 1.15))

    p = doc.add_paragraph()
    _style_paragraph(p, space_after_pt=14, space_before_pt=0, line_spacing=line_spacing)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), accent_hex.lstrip("#").upper())
    pBdr.append(bottom)
    pPr.append(pBdr)
    run = p.add_run("")
    run.font.size = Pt(1)


def _parse_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _style_table(table, theme: dict) -> None:
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    colors = _theme_colors(theme)
    typo = _theme_type(theme)
    border_hex = (theme.get("colors") or {}).get("soft_border") or "#E5E5E5"
    body_font = typo["body_font"]

    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            _set_cell_borders(cell, color_hex=border_hex, sz="4")
            if r_idx == 0:
                _shade_cell(cell, colors["accent"])
            elif r_idx % 2 == 0:
                _shade_cell(cell, colors["alt_row"])

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    if r_idx == 0:
                        _set_run_font(
                            run,
                            size_pt=int(typo["table_header_size_pt"]),
                            bold=True,
                            color=colors["white"],
                            font=body_font,
                        )
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    else:
                        _set_run_font(
                            run,
                            size_pt=int(typo["table_body_size_pt"]),
                            bold=False,
                            color=colors["near_black"],
                            font=body_font,
                        )


def _add_markdown_content(doc, content: str, theme: dict) -> None:
    colors = _theme_colors(theme)
    typo = _theme_type(theme)
    layout = _theme_layout(theme)
    line_spacing = float(layout.get("line_spacing", 1.15))
    body_font = typo["body_font"]
    heading_font = typo["heading_font"]
    accent = colors["accent"]
    near_black = colors["near_black"]
    muted = colors["muted_gray"]

    heading_sizes = {
        1: int(typo["h1_size_pt"]),
        2: int(typo["h2_size_pt"]),
        3: int(typo["h3_size_pt"]),
    }

    lines = content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if _HR_RE.match(stripped):
            if layout.get("title_rule", True):
                _add_brand_rule(doc, theme)
            i += 1
            continue

        heading = _HEADING_RE.match(stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            h = doc.add_heading(heading.group(2).strip(), level=level)
            for run in h.runs:
                _set_run_font(
                    run,
                    size_pt=heading_sizes[level],
                    bold=True,
                    color=accent,
                    font=heading_font,
                )
            i += 1
            continue

        bullet = _BULLET_RE.match(stripped)
        if bullet:
            p = doc.add_paragraph(bullet.group(1).strip(), style="List Bullet")
            _style_paragraph(p, space_after_pt=4, line_spacing=line_spacing)
            for run in p.runs:
                _set_run_font(
                    run,
                    size_pt=int(typo["body_size_pt"]),
                    color=near_black,
                    font=body_font,
                )
            i += 1
            continue

        numbered = _NUMBERED_RE.match(stripped)
        if numbered:
            p = doc.add_paragraph(numbered.group(1).strip(), style="List Number")
            _style_paragraph(p, space_after_pt=4, line_spacing=line_spacing)
            for run in p.runs:
                _set_run_font(
                    run,
                    size_pt=int(typo["body_size_pt"]),
                    color=near_black,
                    font=body_font,
                )
            i += 1
            continue

        if "|" in stripped and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1].strip()):
            headers = _parse_table_row(stripped)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and "|" in lines[i]:
                row_line = lines[i].strip()
                if not row_line or _TABLE_SEP_RE.match(row_line):
                    i += 1
                    continue
                rows.append(_parse_table_row(row_line))
                i += 1
            cols = max(len(headers), max((len(r) for r in rows), default=0))
            if cols > 0:
                table = doc.add_table(rows=1 + len(rows), cols=cols)
                table.style = "Table Grid"
                for c, text in enumerate(headers):
                    if c < cols:
                        table.rows[0].cells[c].text = text
                for r_idx, row in enumerate(rows):
                    for c, text in enumerate(row):
                        if c < cols:
                            table.rows[r_idx + 1].cells[c].text = text
                _style_table(table, theme)
                doc.add_paragraph("")
            continue

        para = doc.add_paragraph()
        _style_paragraph(para, line_spacing=line_spacing)
        run = para.add_run(stripped)
        if stripped.lower().startswith("note:") or stripped.startswith("_"):
            _set_run_font(
                run,
                size_pt=int(typo["notes_size_pt"]),
                color=muted,
                font=body_font,
            )
        else:
            _set_run_font(
                run,
                size_pt=int(typo["body_size_pt"]),
                color=near_black,
                font=body_font,
            )
        i += 1


def create_docx_document(
    title: str,
    content: str,
    filename: Optional[str] = None,
    skill_name: Optional[str] = None,
) -> str:
    """Create a .docx from title + markdown-ish body and return a download link.

    Visual theme is loaded from ``skills/<skill_name>/references/theme.yml``
    (default skill: ``create-docx-report``). Report structure should follow that
    skill's ``SKILL.md``.

    Args:
        title: Document title (Heading 1).
        content: Body text (light markdown).
        filename: Optional download name.
        skill_name: Skill folder whose theme.yml to apply.

    Returns:
        Message including a markdown download link for the chat UI.
    """
    from docx import Document

    if not (title or "").strip() and not (content or "").strip():
        return "Error: provide a non-empty title and/or content for the document."

    skill = (skill_name or _DEFAULT_SKILL).strip() or _DEFAULT_SKILL
    theme = load_skill_theme(skill)
    colors = _theme_colors(theme)
    typo = _theme_type(theme)
    layout = _theme_layout(theme)

    doc = Document()
    _apply_document_defaults(doc, theme)

    doc_title = (title or "Report").strip()
    heading = doc.add_heading(doc_title, level=1)
    for run in heading.runs:
        _set_run_font(
            run,
            size_pt=int(typo["title_size_pt"]),
            bold=True,
            color=colors["accent"],
            font=typo["heading_font"],
        )
    if layout.get("title_rule", True):
        _add_brand_rule(doc, theme)

    if content and content.strip():
        _add_markdown_content(doc, content.strip(), theme)

    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()

    name = (filename or "").strip() or f"{_slugify(doc_title)}.docx"
    meta = save_artifact(data, filename=name, media_type=DOCX_MEDIA_TYPE)
    path = meta["download_path"]
    fname = meta["filename"]

    logger.info(
        "Created DOCX artifact %s (%s, %d bytes) theme_skill=%s",
        meta["id"],
        fname,
        meta["size"],
        skill,
    )
    return (
        f"Document created successfully.\n"
        f"Filename: {fname}\n"
        f"Theme skill: {skill}\n"
        f"Download path: {path}\n\n"
        f"Include this exact markdown link in your reply so the user can download the file:\n"
        f"[{fname}]({path})"
    )


def _slugify(text: str) -> str:
    slug = re.sub(r"[^\w\-]+", "_", text.strip(), flags=re.UNICODE).strip("_")
    return (slug[:60] or "document").lower()


def create_docx_tool() -> StructuredTool:
    """Return the create_docx LangChain tool for the supervisor."""
    return StructuredTool.from_function(
        func=create_docx_document,
        name="create_docx",
        description=(
            "Create a Microsoft Word (.docx) document the user can download. "
            "Visual styling is loaded from a skill's references/theme.yml "
            f"(default skill_name='{_DEFAULT_SKILL}'). Follow that skill's "
            "SKILL.md for section structure. Pass title, content (light "
            "markdown: headings, bullets, numbered lists, pipe tables, --- "
            "rules), optional filename, and optional skill_name. Returns a "
            "optional skill_name. Returns a markdown download link that MUST "
            "be included verbatim in your final reply."
        ),
    )
